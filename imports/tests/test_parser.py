import csv
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from django.test import SimpleTestCase
from openpyxl import Workbook
from imports.parser import CsvImportError, parse_cashbook_csv, parse_cashbook_xlsx

FIXTURE = Path(__file__).parent / "fixtures" / "cashbook_july_synthetic.csv"


class CashbookParserTests(SimpleTestCase):
    def content(self):
        return FIXTURE.read_bytes()

    def test_realistic_metadata_footer_money_dates_and_normalization(self):
        parsed = parse_cashbook_csv(self.content())
        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.period, date(2026, 7, 1))
        self.assertEqual(len(parsed.valid_rows), 19)
        self.assertEqual(sum(row.classification == "TAXABLE" for row in parsed.valid_rows), 13)
        self.assertEqual(sum(row.classification == "EXEMPT" for row in parsed.valid_rows), 6)
        self.assertEqual(parsed.totals, {"net": 546034, "vat": 72966, "total": 619000})
        self.assertEqual(
            (parsed.declared_net, parsed.declared_vat, parsed.declared_total),
            (546034, 72966, 619000),
        )
        self.assertEqual(parsed.valid_rows[0].normalized["issue_date"], "2026-07-01")
        self.assertEqual(parsed.valid_rows[0].normalized["money_layout"], "HISTORICAL_INVERTED")
        self.assertEqual(parsed.valid_rows[-1].normalized["money_layout"], "CANONICAL")
        self.assertEqual(sum(row.status == "IGNORED" for row in parsed.rows), 7)

    def test_utf8_bom_is_supported(self):
        parsed = parse_cashbook_csv(b"\xef\xbb\xbf" + self.content())
        self.assertEqual(len(parsed.valid_rows), 19)

    def test_csv_delimiter_fallback_handles_variable_width_metadata(self):
        with patch("imports.parser.csv.Sniffer.sniff", side_effect=csv.Error):
            parsed = parse_cashbook_csv(self.content())
        self.assertEqual(parsed.errors, [])
        self.assertEqual(len(parsed.valid_rows), 19)

    def test_taxable_canonical_money_layout_is_supported(self):
        content = self.content().replace(
            b'"$35,000","$5,588","$29,412"',
            b'"$29,412","$5,588","$35,000"',
            1,
        )
        parsed = parse_cashbook_csv(content)
        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.valid_rows[0].normalized["money_layout"], "CANONICAL")
        self.assertEqual(parsed.totals, {"net": 546034, "vat": 72966, "total": 619000})

    def test_size_and_row_limits_are_enforced(self):
        with self.assertRaises(CsvImportError):
            parse_cashbook_csv(b"a" * (2 * 1024 * 1024 + 1))
        many_rows = b"a,b\n" * 501
        with self.assertRaises(CsvImportError):
            parse_cashbook_csv(many_rows)

    def test_inconsistent_footer_blocks(self):
        content = self.content().replace(b'"$619,000"', b'"$620,000"')
        parsed = parse_cashbook_csv(content)
        self.assertIn("Totales declarados inválidos", " ".join(parsed.errors))

    def test_invalid_encoding_and_binary_content_are_rejected(self):
        with self.assertRaises(CsvImportError):
            parse_cashbook_csv(b"\xff\xfeinvalid")
        with self.assertRaises(CsvImportError):
            parse_cashbook_csv(b"a,b\x00c")

    def test_formula_prefix_is_flagged_without_changing_original(self):
        content = self.content().replace(b"Servicio sint\xc3\xa9tico 01", b"=2+2", 1)
        parsed = parse_cashbook_csv(content)
        row = parsed.valid_rows[0]
        self.assertEqual(row.normalized["detail"], "=2+2")
        self.assertIn("prefijo de fórmula", " ".join(row.warnings))

    def test_xlsx_selects_period_and_normalizes_historical_layout(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Julio 2026"
        sheet["B2"] = "Rut"
        sheet["C2"] = "76.555.555-5"
        sheet["B3"] = "Razón Social"
        sheet["C3"] = "Organización Sintética SpA"
        sheet["B4"] = "Periodo"
        sheet["C4"] = date(2026, 7, 1)
        headers = [
            "N° Correlativo", "Tipo de Operación", "N° de Documento",
            "Tipo Documento", "Rut Emisor", "Afecta / No Afecta",
            "Fecha de Operación", "Detalle", "Monto Neto", "Monto Iva", "Monto Total",
        ]
        for column, header in enumerate(headers, start=2):
            sheet.cell(6, column, header)
        values = [1, "Ingreso", 20001, "Boleta Electrónica", "11111111-1", "Afecta", date(2026, 7, 20), "Clase", 36000, 5747.9, 30252.1]
        for column, value in enumerate(values, start=2):
            sheet.cell(7, column, value)
        for column, value in zip((9, 10, 11, 12), ("Total", 36000, 5747.9, 30252.1)):
            sheet.cell(8, column, value)
        output = BytesIO()
        workbook.save(output)

        parsed = parse_cashbook_xlsx(output.getvalue(), expected_period=date(2026, 7, 1))

        self.assertEqual(parsed.errors, [])
        self.assertEqual(len(parsed.valid_rows), 1)
        self.assertEqual(
            (
                parsed.valid_rows[0].normalized["net_amount"],
                parsed.valid_rows[0].normalized["vat_amount"],
                parsed.valid_rows[0].normalized["total_amount"],
            ),
            (30252, 5748, 36000),
        )

    def test_xlsx_exempt_row_ignores_stray_vat_and_combines_split_detail(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Marzo 2026"
        sheet.append([])
        sheet.append([None, "Rut", "76.555.555-5"])
        sheet.append([None, "Razón Social", "Organización Sintética SpA"])
        sheet.append([None, "Periodo", date(2026, 3, 1)])
        sheet.append([])
        sheet.append([None, "N° Correlativo", "Tipo de Operación", "N° de Documento", "Tipo Documento", "Rut Emisor", "Afecta / No Afecta", "Fecha de Operación", "Column 12", "Detalle", "Monto Neto", "Monto Iva", "Monto Total"])
        sheet.append([None, 1, "Ingreso", 25040, "Boleta Electrónica", "11111111-1", "Exenta", date(2026, 3, 21), "Taller -", "Persona", 21600, 1436, 20164])
        sheet.append([None, None, None, None, None, None, None, None, "Total", None, 21600, 1436, 20164])
        output = BytesIO()
        workbook.save(output)

        parsed = parse_cashbook_xlsx(output.getvalue(), expected_period=date(2026, 3, 1))

        row = parsed.valid_rows[0]
        self.assertEqual(row.normalized["detail"], "Taller - Persona")
        self.assertEqual((row.normalized["net_amount"], row.normalized["vat_amount"], row.normalized["total_amount"]), (21600, 0, 21600))
        self.assertIn("IVA cero", " ".join(row.warnings))
