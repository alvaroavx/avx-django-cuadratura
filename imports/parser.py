import csv
import hashlib
import io
import json
import re
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

PARSER_VERSION = "cashbook_csv_v1"
XLSX_PARSER_VERSION = "cashbook_xlsx_v1"
MAX_FILE_SIZE = 2 * 1024 * 1024
MAX_XLSX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 500
MAX_XLSX_SHEET_ROWS = 5000


class CsvImportError(ValueError):
    pass


def normalize_text(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip())
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def normalize_tax_id(value):
    return re.sub(r"[^0-9Kk]", "", str(value or "")).upper()


def parse_clp(value):
    raw = str(value or "").strip()
    if not raw:
        return 0
    cleaned = raw.replace("$", "").replace(" ", "").replace(".", "")
    if cleaned.endswith(",00"):
        cleaned = cleaned[:-3]
    else:
        cleaned = cleaned.replace(",", "")
    if not re.fullmatch(r"\d+", cleaned):
        raise CsvImportError(f"Monto CLP inválido: {raw}")
    return int(cleaned)


def parse_xlsx_clp(value):
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        raise CsvImportError(f"Monto CLP inválido: {value}")
    if isinstance(value, (int, float, Decimal)):
        try:
            return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        except InvalidOperation as exc:
            raise CsvImportError(f"Monto CLP inválido: {value}") from exc
    return parse_clp(value)


def parse_chilean_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y-%m"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise CsvImportError(f"Fecha inválida: {raw}")


def _normalize_money(net_header, vat_header, total_header, classification):
    source_net = parse_clp(net_header)
    source_vat = parse_clp(vat_header)
    source_total = parse_clp(total_header)
    if classification == "EXEMPT":
        if source_vat != 0:
            raise CsvImportError("Una fila exenta debe tener IVA cero.")
        if source_net != source_total:
            raise CsvImportError("El monto exento debe coincidir en neto y total.")
        return source_net, 0, source_total, "CANONICAL"
    if source_net + source_vat == source_total:
        return source_net, source_vat, source_total, "CANONICAL"
    if source_net - source_vat == source_total:
        return source_total, source_vat, source_net, "HISTORICAL_INVERTED"
    raise CsvImportError("Los montos no cumplen neto + IVA = total ni el formato histórico invertido.")


@dataclass
class ParsedRow:
    line_number: int
    original: dict
    normalized: dict = field(default_factory=dict)
    fingerprint: str = ""
    classification: str = ""
    warnings: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    status: str = "NEW"


@dataclass
class ParsedImport:
    file_hash: str
    file_size: int
    taxpayer_tax_id: str = ""
    taxpayer_name: str = ""
    period: date | None = None
    header_line: int = 0
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    declared_net: int | None = None
    declared_vat: int | None = None
    declared_total: int | None = None

    @property
    def operation_rows(self):
        return [row for row in self.rows if row.status != "IGNORED"]

    @property
    def valid_rows(self):
        return [row for row in self.rows if row.status == "NEW"]

    @property
    def totals(self):
        valid = [row for row in self.rows if row.status not in {"IGNORED", "ERROR"} and row.normalized]
        return {
            "net": sum(row.normalized["net_amount"] for row in valid),
            "vat": sum(row.normalized["vat_amount"] for row in valid),
            "total": sum(row.normalized["total_amount"] for row in valid),
        }


HEADER_ALIASES = {
    "n correlativo": "correlative",
    "tipo de operacion": "operation_type",
    "n de documento": "folio",
    "tipo documento": "document_name",
    "rut emisor": "historical_issuer",
    "afecta no afecta": "classification",
    "fecha de operacion": "issue_date",
    "detalle": "detail",
    "glosa de operacion": "detail",
    "monto neto": "source_net",
    "monto iva": "source_vat",
    "monto total": "source_total",
}


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _xlsx_original_dict(headers, values):
    result = {}
    for index, value in enumerate(values):
        header = str(headers[index] or "").strip() if index < len(headers) else ""
        result[header or f"_column_{index + 1}"] = _json_value(value)
    return result


def _xlsx_money(source_net, source_vat, source_total, classification):
    net_value = parse_xlsx_clp(source_net)
    vat_value = parse_xlsx_clp(source_vat)
    total_value = parse_xlsx_clp(source_total)
    warnings = []
    if classification == "EXEMPT":
        gross = max(net_value, total_value)
        if gross <= 0:
            raise CsvImportError("El total exento debe ser mayor que cero.")
        if vat_value:
            warnings.append("La fila exenta contenía IVA en la planilla; se normalizó a IVA cero.")
        return gross, 0, gross, "EXEMPT_NORMALIZED", warnings
    net, vat, total, mode = _normalize_money(
        str(net_value), str(vat_value), str(total_value), classification
    )
    return net, vat, total, mode, warnings


def _xlsx_metadata(sheet, label):
    wanted = normalize_text(label)
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True):
        for index, value in enumerate(row):
            if normalize_text(value) == wanted:
                for candidate in row[index + 1 :]:
                    if candidate not in (None, ""):
                        return candidate
    return None


def _xlsx_header(sheet):
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), values_only=True), start=1
    ):
        candidate = {
            HEADER_ALIASES.get(normalize_text(value)): position
            for position, value in enumerate(row)
            if HEADER_ALIASES.get(normalize_text(value))
        }
        required = (
            "correlative", "folio", "classification", "issue_date", "detail",
            "source_net", "source_vat", "source_total",
        )
        if all(key in candidate for key in required):
            return row_number, list(row), candidate
    return None, [], {}


def parse_cashbook_xlsx(content: bytes, *, expected_period: date):
    if not content:
        raise CsvImportError("El archivo está vacío.")
    if len(content) > MAX_XLSX_FILE_SIZE:
        raise CsvImportError("El archivo XLSX supera el límite de 10 MiB.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise CsvImportError("No se pudo leer el archivo XLSX.") from exc
    parsed = ParsedImport(file_hash=hashlib.sha256(content).hexdigest(), file_size=len(content))
    expected_period = expected_period.replace(day=1)
    candidates = []
    for sheet in workbook.worksheets:
        period_value = _xlsx_metadata(sheet, "Periodo") or _xlsx_metadata(sheet, "Período")
        try:
            sheet_period = parse_chilean_date(period_value).replace(day=1)
        except CsvImportError:
            continue
        if sheet_period == expected_period:
            candidates.append(sheet)
    if not candidates:
        parsed.errors.append(f"No se encontró una hoja para el período {expected_period:%Y-%m}.")
        return parsed
    if len(candidates) > 1:
        parsed.errors.append(f"Hay más de una hoja para el período {expected_period:%Y-%m}.")
        return parsed
    sheet = candidates[0]
    parsed.period = expected_period
    parsed.taxpayer_tax_id = str(_xlsx_metadata(sheet, "RUT") or "").strip()
    parsed.taxpayer_name = str(_xlsx_metadata(sheet, "Razón Social") or "").strip()
    header_line, headers, header_map = _xlsx_header(sheet)
    if not header_line:
        parsed.errors.append("No se encontró la cabecera versionada del libro de caja.")
        return parsed
    parsed.header_line = header_line
    if sheet.max_row - header_line > MAX_XLSX_SHEET_ROWS:
        raise CsvImportError(
            f"La hoja supera el límite técnico de {MAX_XLSX_SHEET_ROWS} filas."
        )
    footer_found = False
    for line_number in range(header_line + 1, sheet.max_row + 1):
        values = [sheet.cell(line_number, column + 1).value for column in range(sheet.max_column)]
        original = _xlsx_original_dict(headers, values)
        normalized_values = {normalize_text(value) for value in values if value not in (None, "")}
        footer_labels = " ".join(normalized_values)
        if any(label in footer_labels for label in ("iva debito fiscal", "retencion segunda categoria", "ppm")):
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=["Campo tributario mensual, no es una operación."]))
            continue
        folio_raw = values[header_map["folio"]] if header_map["folio"] < len(values) else None
        correlative = values[header_map["correlative"]] if header_map["correlative"] < len(values) else None
        if "total" in normalized_values:
            footer_found = True
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=["Fila de totales."]))
            continue
        if folio_raw in (None, "", 0, 0.0, "0"):
            reason = "Fila vacía o de plantilla." if not any(value not in (None, "", 0, 0.0) for value in values) or correlative else "Fila de plantilla sin documento."
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=[reason]))
            continue
        operation_index = header_map.get("operation_type")
        if operation_index is not None and normalize_text(values[operation_index]) != "ingreso":
            parsed.rows.append(
                ParsedRow(
                    line_number,
                    original,
                    status="IGNORED",
                    warnings=["Fila auxiliar, no es una operación de ingreso."],
                )
            )
            continue
        row_result = ParsedRow(line_number, original)
        try:
            folio_decimal = Decimal(str(folio_raw))
            if folio_decimal != folio_decimal.to_integral_value() or folio_decimal <= 0:
                raise CsvImportError("El folio debe ser un entero positivo.")
            folio = int(folio_decimal)
            classification_text = normalize_text(values[header_map["classification"]])
            if classification_text == "afecta":
                classification, document_type = "TAXABLE", 39
            elif classification_text in {"exenta", "no afecta", "no afecta o exenta"}:
                classification, document_type = "EXEMPT", 41
            else:
                raise CsvImportError("Clasificación afecta/exenta desconocida.")
            issue_date = parse_chilean_date(values[header_map["issue_date"]])
            detail_parts = []
            detail_index = header_map["detail"]
            if detail_index > 0 and normalize_text(headers[detail_index - 1]).startswith("column"):
                detail_parts.append(str(values[detail_index - 1] or "").strip())
            detail_parts.append(str(values[detail_index] or "").strip())
            detail = " ".join(part for part in detail_parts if part).strip()
            net, vat, total, mode, money_warnings = _xlsx_money(
                values[header_map["source_net"]],
                values[header_map["source_vat"]],
                values[header_map["source_total"]],
                classification,
            )
            if not detail:
                raise CsvImportError("El detalle histórico está vacío.")
            unnamed_values = [
                str(value).strip()
                for index, value in enumerate(values)
                if value not in (None, "")
                and (index >= len(headers) or not str(headers[index] or "").strip())
            ]
            observation = unnamed_values[-1] if unnamed_values else ""
            normalized = {
                "correlative": str(correlative or ""),
                "folio": folio,
                "document_type": document_type,
                "classification": classification,
                "issue_date": issue_date.isoformat(),
                "detail": detail,
                "historical_observation": observation,
                "historical_issuer": str(values[header_map.get("historical_issuer", -1)] or "").strip() if header_map.get("historical_issuer") is not None else "",
                "net_amount": net,
                "vat_amount": vat,
                "total_amount": total,
                "money_layout": mode,
                "source_sheet": sheet.title,
            }
            row_result.classification = classification
            row_result.normalized = normalized
            row_result.warnings.extend(money_warnings)
            if mode == "HISTORICAL_INVERTED":
                row_result.warnings.append("Columnas Monto Neto y Monto Total históricamente invertidas; valores normalizados.")
            row_result.warnings.append("Pago no informado: no se creará un Payment.")
            if detail.lstrip().startswith(("=", "+", "-", "@")):
                row_result.warnings.append("Detalle con prefijo de fórmula; se neutralizará al exportar.")
            fingerprint_payload = {
                key: normalized[key]
                for key in ("folio", "document_type", "classification", "issue_date", "detail", "net_amount", "vat_amount", "total_amount")
            }
            row_result.fingerprint = hashlib.sha256(
                json.dumps(fingerprint_payload, sort_keys=True, ensure_ascii=False).encode()
            ).hexdigest()
        except (CsvImportError, InvalidOperation, ValueError) as exc:
            row_result.status = "ERROR"
            row_result.errors.append(str(exc))
        parsed.rows.append(row_result)
    if len(parsed.operation_rows) > MAX_ROWS:
        raise CsvImportError(f"La hoja supera el límite de {MAX_ROWS} operaciones.")
    if not footer_found:
        parsed.warnings.append("La hoja no contiene una fila de totales; se usarán los totales calculados.")
    return parsed


def _metadata_value(rows, label):
    wanted = normalize_text(label)
    for row in rows:
        for index, value in enumerate(row):
            if normalize_text(value) == wanted:
                for candidate in row[index + 1 :]:
                    if str(candidate).strip():
                        return str(candidate).strip()
    return ""


def _original_dict(headers, row):
    result = {}
    for index, value in enumerate(row):
        header = headers[index].strip() if index < len(headers) else ""
        result[header or f"_column_{index + 1}"] = value
    return result


def parse_cashbook_csv(content: bytes):
    if not content:
        raise CsvImportError("El archivo está vacío.")
    if len(content) > MAX_FILE_SIZE:
        raise CsvImportError("El archivo supera el límite de 2 MiB.")
    if b"\x00" in content:
        raise CsvImportError("El archivo no parece ser un CSV de texto.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CsvImportError("El archivo debe estar codificado en UTF-8.") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    except csv.Error:
        delimiter = ";" if text[:4096].count(";") > text[:4096].count(",") else ","
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    if len(rows) > MAX_ROWS:
        raise CsvImportError(f"El archivo supera el límite de {MAX_ROWS} filas.")
    parsed = ParsedImport(file_hash=hashlib.sha256(content).hexdigest(), file_size=len(content))
    header_index = None
    header_map = {}
    for index, row in enumerate(rows):
        candidate = {HEADER_ALIASES.get(normalize_text(value)): pos for pos, value in enumerate(row)}
        if all(candidate.get(key) is not None for key in ("correlative", "folio", "classification", "issue_date", "detail", "source_net", "source_vat", "source_total")):
            header_index = index
            header_map = candidate
            break
    if header_index is None:
        parsed.errors.append("No se encontró la cabecera versionada del libro de caja.")
        return parsed
    parsed.header_line = header_index + 1
    metadata = rows[:header_index]
    parsed.taxpayer_tax_id = _metadata_value(metadata, "RUT")
    parsed.taxpayer_name = _metadata_value(metadata, "Razón Social")
    period_value = _metadata_value(metadata, "Periodo") or _metadata_value(metadata, "Período")
    try:
        period_date = parse_chilean_date(period_value)
        parsed.period = period_date.replace(day=1)
    except CsvImportError:
        parsed.errors.append("El período de la cabecera es inválido o está ausente.")

    headers = rows[header_index]
    footer_found = False
    for line_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        original = _original_dict(headers, row)
        values = [str(value).strip() for value in row]
        normalized_values = {normalize_text(value) for value in values if value}
        correlative = values[header_map["correlative"]] if header_map["correlative"] < len(values) else ""
        folio_raw = values[header_map["folio"]] if header_map["folio"] < len(values) else ""
        detail_raw = values[header_map["detail"]] if header_map["detail"] < len(values) else ""
        footer_labels = " ".join(normalized_values)
        if any(label in footer_labels for label in ("iva debito fiscal", "retencion segunda categoria", "ppm")):
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=["Campo tributario mensual, no es una operación."]))
            continue
        if "total" in normalized_values and not folio_raw:
            footer_found = True
            try:
                net, vat, total, mode = _normalize_money(
                    values[header_map["source_net"]], values[header_map["source_vat"]],
                    values[header_map["source_total"]], "TAXABLE"
                )
                parsed.declared_net, parsed.declared_vat, parsed.declared_total = net, vat, total
                if mode == "HISTORICAL_INVERTED":
                    parsed.warnings.append("Los totales declarados usan columnas monetarias históricas invertidas.")
            except CsvImportError as exc:
                parsed.errors.append(f"Totales declarados inválidos en línea {line_number}: {exc}")
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=["Fila de totales."]))
            continue
        if not folio_raw:
            reason = "Fila vacía o de plantilla." if not any(values) or correlative else "Correlativo sin documento."
            parsed.rows.append(ParsedRow(line_number, original, status="IGNORED", warnings=[reason]))
            continue
        row_result = ParsedRow(line_number, original)
        try:
            folio = int(re.sub(r"\D", "", folio_raw))
            classification_text = normalize_text(values[header_map["classification"]])
            if classification_text == "afecta":
                classification, document_type = "TAXABLE", 39
            elif classification_text in {"exenta", "no afecta", "no afecta o exenta"}:
                classification, document_type = "EXEMPT", 41
            else:
                raise CsvImportError("Clasificación afecta/exenta desconocida.")
            issue_date = parse_chilean_date(values[header_map["issue_date"]])
            net, vat, total, mode = _normalize_money(
                values[header_map["source_net"]], values[header_map["source_vat"]],
                values[header_map["source_total"]], classification
            )
            unnamed = [value.strip() for index, value in enumerate(values) if index >= len(headers) or not headers[index].strip()]
            observation = next((value for value in reversed(unnamed) if value), "")
            normalized = {
                "correlative": correlative,
                "folio": folio,
                "document_type": document_type,
                "classification": classification,
                "issue_date": issue_date.isoformat(),
                "detail": detail_raw.strip(),
                "historical_observation": observation,
                "historical_issuer": values[header_map["historical_issuer"]].strip() if header_map.get("historical_issuer") is not None else "",
                "net_amount": net,
                "vat_amount": vat,
                "total_amount": total,
                "money_layout": mode,
            }
            if not normalized["detail"]:
                raise CsvImportError("El detalle histórico está vacío.")
            row_result.classification = classification
            row_result.normalized = normalized
            if mode == "HISTORICAL_INVERTED":
                row_result.warnings.append("Columnas Monto Neto y Monto Total históricamente invertidas; valores normalizados.")
            row_result.warnings.append("Pago no informado: no se creará un Payment.")
            if detail_raw.lstrip().startswith(("=", "+", "-", "@")):
                row_result.warnings.append("Detalle con prefijo de fórmula; se neutralizará al exportar.")
            fingerprint_payload = {key: normalized[key] for key in ("folio", "document_type", "classification", "issue_date", "detail", "net_amount", "vat_amount", "total_amount")}
            row_result.fingerprint = hashlib.sha256(json.dumps(fingerprint_payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        except (CsvImportError, ValueError) as exc:
            row_result.status = "ERROR"
            row_result.errors.append(str(exc))
        parsed.rows.append(row_result)
    if not footer_found:
        parsed.errors.append("No se encontró una fila de totales declarados.")
    totals = parsed.totals
    if None not in (parsed.declared_net, parsed.declared_vat, parsed.declared_total):
        if (totals["net"], totals["vat"], totals["total"]) != (parsed.declared_net, parsed.declared_vat, parsed.declared_total):
            parsed.errors.append("Los totales calculados no coinciden con los totales declarados.")
    return parsed
