import hashlib
from datetime import date
from io import BytesIO
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from accounts.models import User
from audit.models import AuditEvent
from documents.models import ElectronicReceipt
from organizations.models import OrganizationMembership, TaxpayerOrganization
from sales.models import Sale, SaleLine
from sales.services import register_manual_operation, register_manual_receipt
from sii.gateway import DisabledSiiGateway, SiiDisabledError
from imports.parser import parse_cashbook_xlsx
from .services import (
    HEADERS,
    close_month,
    generate_receipt_export_workbook,
    generate_workbook,
    prepare_receipt_export,
    reopen_month,
)

class ReportingTests(TestCase):
    def setUp(self):
        self.user=User.objects.create_user("accountant")
        self.org=TaxpayerOrganization.objects.create(tax_id="76.200.001-1",legal_name="Reporte SpA",dte_39_enabled=True)
        sale,_=register_manual_operation(organization=self.org,actor=self.user,operation_date=date(2026,7,3),payment_date=date(2026,7,3),total=1190,transfer_reference="R",detail="Actividad",classification="TAXABLE",idempotency_key=__import__("uuid").uuid4())
        register_manual_receipt(organization=self.org,actor=self.user,sale=sale,document_type=39,folio=15,tax_issue_date=date(2026,7,3),issuer_name="Reporte SpA",signer_name="Persona")

    def test_provisional_export_has_required_sheets_totals_and_manual_fields(self):
        content,totals,exceptions=generate_workbook(self.org,date(2026,7,1),withholding=25,ppm=30)
        wb=load_workbook(BytesIO(content),data_only=True)
        self.assertEqual(wb.sheetnames,["Julio 2026","Excepciones"]); self.assertEqual(totals,[1000,190,1190]); self.assertFalse(exceptions)
        values=[cell.value for row in wb["Julio 2026"] for cell in row]
        self.assertIn("Retención segunda categoría",values); self.assertIn("PPM",values)
        self.assertEqual(wb["Julio 2026"]["B6"].value,"N° Correlativo")
        self.assertEqual(wb["Julio 2026"]["L7"].value,1190)
        parsed=parse_cashbook_xlsx(content,expected_period=date(2026,7,1))
        self.assertEqual(parsed.errors,[])
        self.assertEqual(len(parsed.valid_rows),1)
        self.assertEqual(parsed.valid_rows[0].normalized["total_amount"],1190)

    def test_close_hash_audit_and_reopen_reason(self):
        close=close_month(organization=self.org,period=date(2026,7,1),actor=self.user,withholding=25,ppm=30)
        with close.file.open("rb") as stream: self.assertEqual(hashlib.sha256(stream.read()).hexdigest(),close.file_hash)
        self.assertTrue(AuditEvent.objects.filter(action="month.closed",target_id=close.id).exists())
        with self.assertRaises(ValidationError): reopen_month(close=close,actor=self.user,reason="")
        reopen_month(close=close,actor=self.user,reason="Corrección confirmada por contadora")
        self.assertTrue(AuditEvent.objects.filter(action="month.reopened",target_id=close.id).exists())

    def test_second_close_is_rejected(self):
        close_month(organization=self.org,period=date(2026,7,1),actor=self.user)
        with self.assertRaises(ValidationError): close_month(organization=self.org,period=date(2026,7,1),actor=self.user)

    def test_disabled_gateway_never_echoes_identifier_or_secret(self):
        secret="private-secret-value"
        with self.assertRaises(SiiDisabledError) as ctx: DisabledSiiGateway().submit(receipt_id=secret)
        self.assertNotIn(secret,str(ctx.exception)); self.assertIn("deshabilitada",str(ctx.exception))


class ReceiptExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("receipt-exporter")
        self.org = TaxpayerOrganization.objects.create(
            tax_id="76.300.001-2",
            legal_name="Elementos SpA",
            dte_39_enabled=True,
            dte_41_enabled=True,
        )
        OrganizationMembership.objects.create(
            organization=self.org,
            user=self.user,
            role=OrganizationMembership.Role.VIEWER,
        )

    def create_receipt(self, *, folio, detail, classification="TAXABLE", total=1190, day=3):
        sale, _created = register_manual_operation(
            organization=self.org,
            actor=self.user,
            operation_date=date(2026, 7, day),
            payment_date=date(2026, 7, day),
            total=total,
            transfer_reference=f"R-{folio}",
            detail=detail,
            classification=classification,
            idempotency_key=__import__("uuid").uuid4(),
        )
        receipt = register_manual_receipt(
            organization=self.org,
            actor=self.user,
            sale=sale,
            document_type=39 if classification == "TAXABLE" else 41,
            folio=folio,
            tax_issue_date=date(2026, 7, day),
        )
        return sale, receipt

    def test_preparation_maps_exact_columns_and_consolidates_lines_once(self):
        sale, _receipt = self.create_receipt(folio=15, detail="Taller aéreo")
        SaleLine.objects.create(
            sale=sale,
            description="Matrícula\tcon ajuste\nmanual",
            quantity=1,
            unit_amount=0,
            total_amount=0,
        )
        self.create_receipt(
            folio=16,
            detail="Clase exenta",
            classification="EXEMPT",
            total=5000,
            day=4,
        )

        export = prepare_receipt_export(self.org, date(2026, 7, 1))

        self.assertEqual(export["headers"], HEADERS)
        self.assertEqual(len(export["rows"]), 2)
        self.assertEqual(
            export["rows"][0],
            [
                1,
                "Ingreso",
                "15",
                "Boleta electrónica afecta",
                "76.300.001-2",
                "Afecta",
                date(2026, 7, 3),
                "Taller aéreo; Matrícula con ajuste manual",
                1000,
                190,
                1190,
            ],
        )
        self.assertEqual(export["rows"][1][5], "Exenta")
        self.assertEqual(export["rows"][1][8:11], [5000, 0, 5000])

    def test_workbook_matches_preview_types_and_spreadsheet_safety(self):
        self.create_receipt(folio=21, detail="=SUM(A1:A2)")

        content, export = generate_receipt_export_workbook(self.org, date(2026, 7, 1))
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook["Boletas"]

        self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:K2")
        self.assertEqual(sheet["C2"].value, "21")
        self.assertEqual(sheet["C2"].number_format, "@")
        self.assertEqual(sheet["E2"].value, self.org.tax_id)
        self.assertEqual(sheet["E2"].number_format, "@")
        self.assertEqual(sheet["G2"].value.date(), export["rows"][0][6])
        self.assertEqual(sheet["G2"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet["H2"].value, "'=SUM(A1:A2)")
        self.assertIsInstance(sheet["I2"].value, int)
        self.assertEqual(sheet["K2"].value, export["rows"][0][10])

    def test_preview_download_filters_and_tenant_permissions(self):
        self.create_receipt(folio=31, detail="Con tilde: acrobacia aérea")
        self.create_receipt(folio=32, detail="Segunda boleta", day=4)
        other_org = TaxpayerOrganization.objects.create(
            tax_id="76.300.009-8", legal_name="Otra SpA"
        )
        self.client.force_login(self.user)

        preview_url = reverse(
            "receipt-export-preview", args=[self.org.id, "2026-07"]
        )
        response = self.client.get(preview_url)
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["count"], 2)
        self.assertEqual(len(payload["headers"]), 11)
        self.assertEqual(payload["rows"][0][6], "03/07/2026")

        filtered = self.client.get(preview_url, {"exception": "pending"}).json()
        self.assertEqual(filtered["count"], 0)
        self.assertEqual(filtered["filter_label"], "Pendientes")

        download = self.client.get(
            reverse("receipt-export-download", args=[self.org.id, "2026-07"])
        )
        self.assertEqual(download.status_code, 200)
        self.assertIn("Elementos-SpA".lower(), download["Content-Disposition"].lower())
        workbook = load_workbook(BytesIO(download.content), data_only=True)
        self.assertEqual(workbook["Boletas"].max_row, 3)

        forbidden = self.client.get(
            reverse("receipt-export-preview", args=[other_org.id, "2026-07"])
        )
        self.assertEqual(forbidden.status_code, 403)

    def test_nonaccepted_receipt_is_reported_but_not_exported(self):
        sale = Sale.objects.create(
            organization=self.org,
            operation_date=date(2026, 7, 6),
            detail="Boleta anulada",
            tax_classification=Sale.TaxClassification.EXEMPT,
            net_amount=900,
            vat_amount=0,
            total_amount=900,
            status=Sale.Status.CANCELLED,
            origin_system="TEST",
            external_id="cancelled-1",
            created_by=self.user,
        )
        ElectronicReceipt.objects.create(
            organization=self.org,
            sale=sale,
            document_type=41,
            folio=99,
            tax_issue_date=date(2026, 7, 6),
            net_amount=900,
            vat_amount=0,
            total_amount=900,
            status=ElectronicReceipt.Status.CANCELLED,
            created_by=self.user,
        )

        export = prepare_receipt_export(self.org, date(2026, 7, 1))

        self.assertEqual(export["rows"], [])
        self.assertEqual(export["excluded_count"], 1)
