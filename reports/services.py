import hashlib
import re
from io import BytesIO
from django.core.files.base import ContentFile
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Max, Prefetch
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from audit.services import record_event
from documents.models import ElectronicReceipt
from sales.models import Sale, SaleLine
from .models import MonthlyClose

HEADERS = [
    "N° Correlativo",
    "Tipo de Operación",
    "N° de Documento",
    "Tipo Documento",
    "Rut Emisor",
    "Afecta / No Afecta",
    "Fecha de Operación",
    "Detalle",
    "Monto Neto",
    "Monto Iva",
    "Monto Total",
]

MONTH_NAMES = (
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
)

RECEIPT_EXPORT_FILTERS = {
    "pending": "Pendientes",
    "rejected": "Rechazadas",
    "uncertain": "Inciertas",
    "historical_incomplete": "Históricas incompletas",
}

RECEIPT_EXPORT_WIDTHS = (16, 20, 20, 28, 18, 22, 20, 52, 16, 14, 16)

def spreadsheet_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def period_receipts(organization, period):
    return ElectronicReceipt.objects.filter(organization=organization, status=ElectronicReceipt.Status.ACCEPTED, tax_issue_date__year=period.year, tax_issue_date__month=period.month).select_related("sale").order_by("tax_issue_date", "folio")


def normalize_receipt_export_filter(value):
    return value if value in RECEIPT_EXPORT_FILTERS else ""


def receipt_export_queryset(organization, period, exception_filter=""):
    exception_filter = normalize_receipt_export_filter(exception_filter)
    receipts = ElectronicReceipt.objects.filter(
        organization=organization,
        status=ElectronicReceipt.Status.ACCEPTED,
        tax_issue_date__year=period.year,
        tax_issue_date__month=period.month,
    )
    if exception_filter == "pending":
        receipts = receipts.exclude(sale__status=Sale.Status.DOCUMENTED)
    elif exception_filter == "rejected":
        receipts = receipts.filter(status=ElectronicReceipt.Status.REJECTED)
    elif exception_filter == "uncertain":
        receipts = receipts.filter(status=ElectronicReceipt.Status.UNCERTAIN)
    elif exception_filter == "historical_incomplete":
        receipts = receipts.filter(
            sale__reconciliation_status=Sale.ReconciliationStatus.MISSING_PAYMENT_DATA
        )
    return receipts.select_related("sale").prefetch_related(
        Prefetch(
            "sale__lines",
            queryset=SaleLine.objects.order_by("id"),
            to_attr="export_lines",
        )
    ).order_by("tax_issue_date", "document_type", "folio", "id")


def _receipt_detail(receipt):
    sale_detail = re.sub(r"[\t\r\n]+", " ", receipt.sale.detail).strip()
    details = [sale_detail] if sale_detail else []
    for line in getattr(receipt.sale, "export_lines", ()):
        detail = re.sub(r"[\t\r\n]+", " ", line.description).strip()
        if detail and detail not in details:
            details.append(detail)
    return "; ".join(details)


def prepare_receipt_export(organization, period, exception_filter=""):
    exception_filter = normalize_receipt_export_filter(exception_filter)
    rows = []
    for index, receipt in enumerate(
        receipt_export_queryset(organization, period, exception_filter), start=1
    ):
        rows.append(
            [
                index,
                "Ingreso",
                str(receipt.folio),
                receipt.get_document_type_display(),
                organization.tax_id,
                receipt.sale.get_tax_classification_display(),
                receipt.tax_issue_date,
                _receipt_detail(receipt),
                receipt.net_amount,
                receipt.vat_amount,
                receipt.total_amount,
            ]
        )
    excluded = ElectronicReceipt.objects.filter(
        organization=organization,
        tax_issue_date__year=period.year,
        tax_issue_date__month=period.month,
    ).exclude(status=ElectronicReceipt.Status.ACCEPTED).count()
    return {
        "headers": HEADERS,
        "rows": rows,
        "filter": exception_filter,
        "filter_label": RECEIPT_EXPORT_FILTERS.get(exception_filter, "Sin filtro de excepciones"),
        "excluded_count": excluded,
    }


def generate_receipt_export_workbook(organization, period, exception_filter=""):
    export = prepare_receipt_export(organization, period, exception_filter)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Boletas"
    sheet.append(export["headers"])
    for row in export["rows"]:
        sheet.append([spreadsheet_safe(value) for value in row])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 3).number_format = "@"
        sheet.cell(row_number, 5).number_format = "@"
        sheet.cell(row_number, 7).number_format = "dd/mm/yyyy"
        for column in (9, 10, 11):
            sheet.cell(row_number, column).number_format = "#,##0"
    for column, width in enumerate(RECEIPT_EXPORT_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), export


def receipt_export_filename(organization, period):
    organization_name = slugify(organization.trade_name or organization.legal_name) or "organizacion"
    return f"{organization_name}-{period:%Y-%m}-boletas.xlsx"

def generate_workbook(organization, period, withholding=0, ppm=0):
    wb=Workbook(); ws=wb.active; ws.title=f"{MONTH_NAMES[period.month - 1]} {period.year}"
    ws["B2"]="Rut"; ws["C2"]=spreadsheet_safe(organization.tax_id)
    ws["B3"]="Razón Social"; ws["C3"]=spreadsheet_safe(organization.legal_name)
    ws["B4"]="Periodo"; ws["C4"]=period; ws["C4"].number_format="mm/yyyy"
    for column, header in enumerate(HEADERS, start=2): ws.cell(6,column,header)
    receipts=list(period_receipts(organization, period))
    for index, receipt in enumerate(receipts, 1):
        sale=receipt.sale
        row=index+6
        values=[index,"Ingreso",receipt.folio,"Boleta Electrónica",spreadsheet_safe(receipt.signer_name or receipt.issuer_name),sale.get_tax_classification_display(),receipt.tax_issue_date,spreadsheet_safe(sale.detail),receipt.net_amount,receipt.vat_amount,receipt.total_amount]
        for column,value in enumerate(values,start=2): ws.cell(row,column,value)
        ws.cell(row,8).number_format="dd/mm/yyyy"
    totals=[sum(getattr(r, field) for r in receipts) for field in ("net_amount","vat_amount","total_amount")]
    total_row=7+len(receipts)
    ws.cell(total_row,9,"Total")
    for column,value in zip((10,11,12),totals): ws.cell(total_row,column,value)
    ws.cell(total_row+2,3,"IVA débito fiscal"); ws.cell(total_row+2,4,totals[1])
    ws.cell(total_row+3,3,"Retención segunda categoría"); ws.cell(total_row+3,4,withholding)
    ws.cell(total_row+4,3,"PPM"); ws.cell(total_row+4,4,ppm)
    header_fill=PatternFill("solid",fgColor="1F4E78")
    for cell in ws[6]:
        if 2 <= cell.column <= 12:
            cell.fill=header_fill; cell.font=Font(color="FFFFFF",bold=True); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for cell in ws[total_row][8:12]: cell.font=Font(bold=True)
    for column in (10,11,12):
        for row in range(7,total_row+1): ws.cell(row,column).number_format='#,##0'
    widths={2:14,3:18,4:18,5:21,6:16,7:20,8:18,9:64,10:16,11:14,12:16}
    for column,width in widths.items(): ws.column_dimensions[get_column_letter(column)].width=width
    ws.freeze_panes="B7"; ws.auto_filter.ref=f"B6:L{max(6,total_row-1)}"
    ex=wb.create_sheet("Excepciones")
    ex.append(["Estado", "Fecha operación", "Detalle", "Total", "Motivo"])
    exceptions=Sale.objects.filter(organization=organization, operation_date__year=period.year, operation_date__month=period.month).exclude(reconciliation_status=Sale.ReconciliationStatus.RECONCILED)
    for sale in exceptions: ex.append([sale.get_reconciliation_status_display(), sale.operation_date, spreadsheet_safe(sale.detail), sale.total_amount, "No se incluye en libro definitivo mientras esté pendiente"])
    ex.freeze_panes="A2"
    ex.column_dimensions["A"].width=24; ex.column_dimensions["B"].width=18; ex.column_dimensions["C"].width=64; ex.column_dimensions["D"].width=16; ex.column_dimensions["E"].width=54
    out=BytesIO(); wb.save(out); return out.getvalue(), totals, exceptions.exists()

@transaction.atomic
def close_month(*, organization, period, actor, withholding=0, ppm=0):
    if MonthlyClose.objects.filter(organization=organization, period=period, status=MonthlyClose.Status.CLOSED).exists(): raise ValidationError("El período ya está cerrado.")
    content, totals, has_exceptions=generate_workbook(organization, period, withholding, ppm)
    if has_exceptions: raise ValidationError("El período contiene excepciones y no puede cerrarse.")
    version=(MonthlyClose.objects.filter(organization=organization, period=period).aggregate(v=Max("version"))["v"] or 0)+1
    digest=hashlib.sha256(content).hexdigest()
    close=MonthlyClose.objects.create(organization=organization, period=period, version=version, closed_at=timezone.now(), closed_by=actor, file_hash=digest, total_net=totals[0], total_vat=totals[1], total_amount=totals[2], second_category_withholding=withholding, ppm=ppm, status=MonthlyClose.Status.CLOSED)
    close.file.save(f"libro-{organization.id}-{period:%Y-%m}-v{version}.xlsx", ContentFile(content), save=True)
    record_event(organization=organization, actor=actor, action="month.closed", target=close, metadata={"period":period.isoformat(),"version":version,"sha256":digest})
    return close

@transaction.atomic
def reopen_month(*, close, actor, reason):
    if close.status != MonthlyClose.Status.CLOSED: raise ValidationError("Solo se puede reabrir un cierre vigente.")
    if not reason.strip(): raise ValidationError("Debe indicar un motivo.")
    close.status=MonthlyClose.Status.REOPENED; close.reopen_reason=reason.strip(); close.save(update_fields=["status","reopen_reason"])
    record_event(organization=close.organization, actor=actor, action="month.reopened", target=close, metadata={"reason":reason.strip()})
    return close
